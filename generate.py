import argparse
import secrets
import sys
from pathlib import Path

import qrcode
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill

from config import BASE_URL, EXCEL_PATH, QR_DIR, init_db


def load_attendees(excel_path):
    """Read (name, email) pairs regardless of header language (e.g. Arabic) or
    extra/empty columns: the email is the cell containing '@', the name is the
    longest other non-empty cell in the same row."""
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active
    people = []
    seen_emails = set()
    drop = {"\u200a", "\u200b", "\u200e", "\u200f", "\ufeff", "\xa0"}
    for row in ws.iter_rows(values_only=True):
        cells = []
        for value in row:
            if value is None:
                continue
            text = str(value)
            for ch in drop:
                text = text.replace(ch, "")
            text = text.strip()
            if text:
                cells.append(text)
        email = next((c for c in cells if "@" in c), None)
        if not email:
            continue
        email = "".join(email.split()).lower()
        if "@" not in email or email in seen_emails:
            continue
        candidates = [c for c in cells if "@" not in c]
        name = max(candidates, key=len) if candidates else email.split("@")[0]
        seen_emails.add(email)
        people.append((name, email))
    wb.close()
    return people


def generate(excel_path=None):
    excel_path = Path(excel_path) if excel_path else EXCEL_PATH
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")
    people = load_attendees(excel_path)
    if not people:
        raise ValueError("No valid (name, email) rows found in the Excel file.")

    conn = init_db()
    added = 0
    for name, email in people:
        existing = conn.execute(
            "SELECT token FROM attendees WHERE email=?", (email,)
        ).fetchone()
        if existing:
            token = existing["token"]
        else:
            token = secrets.token_urlsafe(12)
            conn.execute(
                "INSERT INTO attendees (name, email, token) VALUES (?, ?, ?)",
                (name, email, token),
            )
            added += 1
        qr_path = QR_DIR / f"{token}.png"
        if not qr_path.exists():
            qrcode.make(f"{BASE_URL}/check/{token}").save(qr_path)
    conn.commit()
    for row in conn.execute("SELECT token FROM attendees"):
        qr_path = QR_DIR / f"{row['token']}.png"
        if not qr_path.exists():
            qrcode.make(f"{BASE_URL}/check/{row['token']}").save(qr_path)

    total = conn.execute("SELECT COUNT(*) AS c FROM attendees").fetchone()["c"]
    export_workbook(conn)
    conn.close()
    return {"total": total, "added": added}


def export_workbook(conn=None, output_path=None):
    output_path = Path(output_path) if output_path else EXCEL_PATH.parent / "attendees_with_qr.xlsx"
    own_conn = conn is None
    if own_conn:
        conn = init_db()
    rows = conn.execute(
        "SELECT name, email, token, scanned FROM attendees ORDER BY id"
    ).fetchall()
    if own_conn:
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendees"
    headers = ["Name", "Email", "QR Code", "Scanned"]
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 12

    for i, row in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=row["name"])
        ws.cell(row=i, column=2, value=row["email"])
        ws.cell(row=i, column=4, value=bool(row["scanned"]))
        qr_path = QR_DIR / f"{row['token']}.png"
        if qr_path.exists():
            img = XLImage(str(qr_path))
            img.width = 90
            img.height = 90
            ws.add_image(img, f"C{i}")
            ws.row_dimensions[i].height = 70

    wb.save(output_path)
    return output_path


def export_status(output_path=None):
    output_path = Path(output_path) if output_path else EXCEL_PATH.parent / "checkin_status.xlsx"
    conn = init_db()
    rows = conn.execute(
        "SELECT name, email, sent, scanned, scanned_at FROM attendees ORDER BY id"
    ).fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Check-in Status"
    headers = ["Name", "Email", "Email Sent", "Scanned", "Scanned At"]
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 20

    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    for i, row in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=row["name"])
        ws.cell(row=i, column=2, value=row["email"])
        ws.cell(row=i, column=3, value=bool(row["sent"]))
        scanned_cell = ws.cell(row=i, column=4, value=bool(row["scanned"]))
        scanned_cell.fill = green if row["scanned"] else red
        ws.cell(row=i, column=5, value=row["scanned_at"] or "")

    wb.save(output_path)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Generate tokens and QR codes from attendees.xlsx")
    parser.add_argument("--excel", default=None, help="Path to attendees Excel file")
    args = parser.parse_args()
    try:
        result = generate(args.excel)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
    print(f"Done. Total in database: {result['total']} | New added: {result['added']}")
    print(f"QR images: {QR_DIR}")
    print(f"Workbook with QRs: attendees_with_qr.xlsx")


if __name__ == "__main__":
    main()
